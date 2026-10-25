"""Streaming, privacy-preserving profiling helpers for uploaded tabular sources."""
from __future__ import annotations

import hashlib
import io
import json
import re
import stat
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Iterable

import pandas as pd

from src.core.config import load_settings
from src.data.profiler import ColumnProfile, DatasetProfileReport


class PartialReadError(ValueError):
    """A parser failed after consuming some rows; counts are exposed for job status."""
    def __init__(self, message: str, processed_rows: int):
        super().__init__(message)
        self.processed_rows = processed_rows


def safe_member_name(name: str) -> bool:
    """Reject absolute, traversal and Windows drive-style archive paths."""
    posix = PurePosixPath(name.replace("\\", "/"))
    return (not posix.is_absolute() and ".." not in posix.parts
            and not re.match(r"^[A-Za-z]:", name) and not name.startswith(("/", "\\")))


def json_record_chunks(stream: BinaryIO, rows: int, chunk_bytes: int = 1024 * 1024):
    """Incrementally decode a top-level JSON array of records from a binary stream."""
    decoder = json.JSONDecoder()
    text = io.TextIOWrapper(stream, encoding="utf-8")
    buffer = ""
    position = 0
    eof = False
    started = False
    batch: list[dict[str, Any]] = []
    while True:
        if not eof:
            block = text.read(chunk_bytes)
            if block:
                buffer = buffer[position:] + block
                position = 0
            else:
                eof = True
        while position < len(buffer) and buffer[position].isspace():
            position += 1
        if not started:
            if position >= len(buffer):
                if eof:
                    raise ValueError("JSON file is empty")
                continue
            if buffer[position] != "[":
                raise ValueError("JSON input must be a top-level array of records; use JSONL for streamed records")
            position += 1
            started = True
        need_more = False
        while position < len(buffer):
            while position < len(buffer) and (buffer[position].isspace() or buffer[position] == ","):
                position += 1
            if position < len(buffer) and buffer[position] == "]":
                position += 1
                if batch:
                    yield pd.DataFrame(batch)
                return
            try:
                value, end = decoder.raw_decode(buffer, position)
            except json.JSONDecodeError as exc:
                if eof:
                    raise ValueError(f"Invalid or incomplete JSON array: {exc.msg}") from exc
                need_more = True
                break
            if not isinstance(value, dict):
                raise ValueError("JSON array entries must be objects")
            batch.append(value)
            position = end
            if len(batch) >= rows:
                yield pd.DataFrame(batch)
                batch = []
        if eof and not need_more:
            raise ValueError("JSON array is missing its closing bracket")


def _chunks(stream: BinaryIO, filename: str, rows: int, archive: zipfile.ZipFile | None = None):
    lower = filename.lower()
    if archive is not None:
        if lower.endswith(".csv"):
            try:
                yield from pd.read_csv(stream, chunksize=rows, low_memory=False)
            except UnicodeDecodeError:
                stream.seek(0)
                yield from pd.read_csv(stream, chunksize=rows, encoding="latin-1", low_memory=False)
            return
        if lower.endswith((".jsonl", ".ndjson")):
            yield from pd.read_json(stream, lines=True, chunksize=rows)
            return
        if lower.endswith(".json"):
            yield from json_record_chunks(stream, rows)
            return
        if lower.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            if hasattr(stream, "seek"):
                stream.seek(0)
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                for sheet_name in workbook.sheetnames:
                    sheet = workbook[sheet_name]
                    iterator = sheet.iter_rows(values_only=True)
                    first_row = next(iterator, None)
                    if not first_row:
                        continue
                    non_empty = [c for c in first_row if c is not None and str(c).strip()]
                    name_lower = sheet_name.strip().lower()
                    if len(non_empty) < 2 or name_lower in {"notes", "note", "readme", "metadata", "instructions", "changelog"}:
                        continue
                    headers = [str(value) if value is not None else f"column_{index + 1}"
                               for index, value in enumerate(first_row)]
                    batch: list[tuple[Any, ...]] = []
                    for row in iterator:
                        batch.append(row)
                        if len(batch) >= rows:
                            yield pd.DataFrame(batch, columns=headers)
                            batch = []
                    if batch:
                        yield pd.DataFrame(batch, columns=headers)
            finally:
                workbook.close()
            return
        if lower.endswith(".xls"):
            raise ValueError("Legacy XLS files are not supported for chunked profiling")
        raise ValueError(f"Unsupported archive member type: {filename}")
    yield from _chunks(stream, filename, rows, archive=object())


def _plain_chunks(stream: BinaryIO, filename: str, rows: int):
    lower = filename.lower()
    if lower.endswith(".csv"):
        try:
            yield from pd.read_csv(stream, chunksize=rows, low_memory=False)
        except UnicodeDecodeError:
            stream.seek(0)
            yield from pd.read_csv(stream, chunksize=rows, encoding="latin-1", low_memory=False)
    elif lower.endswith((".jsonl", ".ndjson")):
        yield from pd.read_json(stream, lines=True, chunksize=rows)
    elif lower.endswith(".json"):
        yield from json_record_chunks(stream, rows)
    elif lower.endswith((".xlsx", ".xlsm")):
        yield from _chunks(stream, filename, rows, archive=object())
    elif lower.endswith(".xls"):
        raise ValueError("Legacy XLS files are not supported for chunked profiling")
    else:
        raise ValueError("Allowed file types: CSV, XLSX, JSON, ZIP")


def profile_uploaded_file(path: Path, filename: str, dataset_name: str, *,
                          chunk_rows: int | None = None) -> DatasetProfileReport:
    """Read data in bounded chunks and retain aggregate statistics only."""
    settings = load_settings().uploads
    chunk_rows = chunk_rows or settings.chunk_rows
    ext = Path(filename).suffix.lower()
    if ext not in {".csv", ".xlsx", ".xlsm", ".json", ".jsonl", ".zip"}:
        raise ValueError("Allowed file types: CSV, XLSX, JSON, ZIP")
    count = 0
    columns: dict[str, dict[str, Any]] = {}
    member_used = None
    formula_risk = False
    sheets_detected: list[str] = []
    sheets_imported: list[str] = []
    sheets_skipped: list[dict[str, str]] = []

    if ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            with path.open("rb") as wb_stream:
                wb = load_workbook(wb_stream, read_only=True, data_only=True)
                sheets_detected = list(wb.sheetnames)
                for s_name in sheets_detected:
                    sheet = wb[s_name]
                    iterator = sheet.iter_rows(values_only=True)
                    first_row = next(iterator, None)
                    if not first_row:
                        sheets_skipped.append({"sheet_name": s_name, "reason": "Empty sheet without data rows"})
                        continue
                    non_empty = [c for c in first_row if c is not None and str(c).strip()]
                    name_lower = s_name.strip().lower()
                    if len(non_empty) < 2 or name_lower in {"notes", "note", "readme", "metadata", "instructions", "changelog"}:
                        sheets_skipped.append({
                            "sheet_name": s_name,
                            "reason": f"Non-tabular sheet (unstructured notes / {len(non_empty)} column(s))",
                        })
                    else:
                        sheets_imported.append(s_name)
                wb.close()
        except Exception:
            pass

    def consume(chunks: Iterable[pd.DataFrame]) -> None:
        nonlocal count, formula_risk
        try:
            for frame in chunks:
                frame.columns = [str(col) for col in frame.columns]
                count += len(frame)
                for name in frame.columns:
                    series = frame[name]
                    state = columns.setdefault(name, {"dtypes": set(), "nulls": 0, "unique": set(),
                        "numeric_n": 0, "numeric_sum": 0.0, "numeric_sq": 0.0,
                        "numeric_min": None, "numeric_max": None, "values": [], "formula": False})
                    state["dtypes"].add(str(series.dtype))
                    state["nulls"] += int(series.isna().sum())
                    vals = series.dropna()
                    for value in vals:
                        token = f"{type(value).__name__}:{value!r}"
                        state["unique"].add(hashlib.sha256(token.encode("utf-8", "replace")).hexdigest())
                    if pd.api.types.is_numeric_dtype(series):
                        nums = pd.to_numeric(vals, errors="coerce").dropna()
                        if len(nums):
                            state["numeric_n"] += len(nums)
                            state["numeric_sum"] += float(nums.sum())
                            state["numeric_sq"] += float((nums * nums).sum())
                            low, high = float(nums.min()), float(nums.max())
                            state["numeric_min"] = low if state["numeric_min"] is None else min(state["numeric_min"], low)
                            state["numeric_max"] = high if state["numeric_max"] is None else max(state["numeric_max"], high)
                    else:
                        text = vals.astype(str)
                        state["formula"] |= bool(text.str.startswith(("=", "+", "-", "@", "\t", "\r")).any())
                        state["values"].extend(text.head(max(0, 500 - len(state["values"]))).tolist())
                        state["values"] = state["values"][:500]
                        formula_risk |= state["formula"]
        except Exception as exc:  # noqa: BLE001 - surfaced as a failed/partial import.
            raise PartialReadError(str(exc), count) from exc

    with path.open("rb") as stream:
        if ext == ".zip":
            with zipfile.ZipFile(stream) as archive:
                infos = archive.infolist()
                if len(infos) > settings.max_zip_entries:
                    raise ValueError(f"Zip archive contains too many entries ({len(infos)} > {settings.max_zip_entries})")
                total_size = sum(info.file_size for info in infos)
                if total_size > settings.max_zip_uncompressed_bytes:
                    raise ValueError("Uncompressed size exceeds configured safety limit")
                compressed = max(path.stat().st_size, 1)
                if total_size / compressed > settings.max_zip_compression_ratio:
                    raise ValueError("Suspicious compression ratio, possible zip bomb")
                for info in infos:
                    if not safe_member_name(info.filename):
                        raise ValueError(f"Path traversal detected in zip entry: '{info.filename}'")
                    if stat.S_ISLNK(info.external_attr >> 16):
                        raise ValueError("Symbolic links are not permitted inside uploaded ZIP files")
                members = [i for i in infos if not i.is_dir() and Path(i.filename).suffix.lower()
                           in {".csv", ".xlsx", ".xlsm", ".json", ".jsonl"}]
                if not members:
                    raise ValueError("ZIP contains no supported CSV, XLSX, or JSON data files")
                members_to_profile = [members[0]] if len(members) > 1 else members
                for info in members_to_profile:
                    member_used = info.filename
                    with archive.open(info) as member:
                        if Path(info.filename).suffix.lower() in {".xlsx", ".xlsm"}:
                            member_bytes = io.BytesIO(member.read())
                            consume(_plain_chunks(member_bytes, info.filename, chunk_rows))
                        else:
                            consume(_plain_chunks(member, info.filename, chunk_rows))
        else:
            consume(_plain_chunks(stream, filename, chunk_rows))
    if not columns:
        raise ValueError("Uploaded dataset contains no rows or columns")

    output: list[ColumnProfile] = []
    for name, state in columns.items():
        unique = len(state["unique"])
        ratio = unique / count if count else 0.0
        vals = pd.Series(state["values"], dtype="string")
        parsed = pd.to_datetime(vals, errors="coerce", format="mixed") if len(vals) else pd.Series([], dtype="datetime64[ns]")
        date_rate = float(parsed.notna().mean()) if len(vals) else 0.0
        formats = set()
        for value in vals.dropna():
            text = str(value)
            has_time = "T" in text or " " in text
            if re.match(r"^\d{4}-\d{1,2}-\d{1,2}", text):
                label = "%Y-%m-%d"
            elif re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", text):
                label = "slash date (day/month order ambiguous)"
            elif re.match(r"^\d{4}/\d{1,2}/\d{1,2}$", text):
                label = "%Y/%m/%d"
            elif re.match(r"^\d{1,2}-\d{1,2}-\d{4}$", text):
                label = "hyphen date (day/month order ambiguous)"
            else:
                label = "other date-like format"
            formats.add(label + (" with time" if has_time else ""))
        numeric_stats = None
        output.append(ColumnProfile(name=name, dtype=" | ".join(sorted(state["dtypes"])), total_rows=count,
            null_count=state["nulls"], null_percentage=round(100 * state["nulls"] / count, 4),
            unique_count=unique, uniqueness_ratio=round(ratio, 4), date_parse_success_rate=round(date_rate, 4),
            date_formats=sorted(formats) if date_rate >= 0.5 else [], candidate_identifier=bool(count and ratio >= 0.95),
            numeric_stats=numeric_stats, formula_injection_risk=state["formula"]))
    return DatasetProfileReport(dataset_name=dataset_name, source_filename=filename, total_rows=count,
        processed_rows=count, total_columns=len(output), columns=output,
        file_type=(f"zip:{member_used}" if member_used else ext.lstrip(".")),
        has_formula_injection_risk=formula_risk, total_rows_known=True,
        sheets_detected=sheets_detected, sheets_imported=sheets_imported, sheets_skipped=sheets_skipped)
